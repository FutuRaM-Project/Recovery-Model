import openpyxl

def import_transfercoefficients_xlsx(filename, model):
    # set the path to the file
    path = filename
    workbook = openpyxl.load_workbook(path, data_only=True)
    # Select the worksheet
    worksheet = workbook.worksheets[0]
    # Get the headers
    headers = [cell.value for cell in worksheet[1]]
    # Convert each row (tc) to a dictionary
    data = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        row_data = {}
        for i, value in enumerate(row):
            if worksheet.cell(row=1, column=i+1).data_type == 's':
                row_data[headers[i]] = str(value)
            elif worksheet.cell(row=1, column=i+1).data_type == 'n':
                row_data[headers[i]] = float(value)
            elif worksheet.cell(row=1, column=i+1).data_type == 'd':
                row_data[headers[i]] = value.date()
            elif worksheet.cell(row=1, column=i+1).data_type == 'b':
                row_data[headers[i]] = bool(value)
            elif worksheet.cell(row=1, column=i+1).data_type == 'f':
                row_data[headers[i]] = worksheet.cell(row=row[0].row, column=row[0].column)._value
            else:
                row_data[headers[i]] = None
        data.append(row_data)

    process_dict = {process.name: process for process in model.processes.values()}

    for tc in data:
        process_name = tc['process']
        if process_name in process_dict:
            process = process_dict[process_name]
            process.add_transfer_coefficient(tc['flow_input'], tc['flow_output'], tc['transfer_coefficient'], tc['uncertainty'])
            # print("Added transfer coefficients to: " + process_name)
        else:
            print("Process not found: " + process_name)

    return data


