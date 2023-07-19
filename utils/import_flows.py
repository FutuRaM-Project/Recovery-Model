import sys

# Add the path to the classes folder to the system path
sys.path.insert(0, '/home/stew/code/gh/futuram/IntegratedModel')


# filename = '/home/stew/code/gh/futuram/ELV/test_system/data/ELV_ICE_flows.xlsx'
# model = Model('test_system_ELV')


import openpyxl
from classes.classes_flows import Flow
from classes.classes_processes import Process
from classes.classes_model import Model

def import_flows_xlsx(filename, model):
    # set the path to the file
    path = filename
    workbook = openpyxl.load_workbook(path, data_only=True)
    # Select the worksheet
    worksheet = workbook.worksheets[0]
    # Get the headers
    headers = [cell.value for cell in worksheet[1]]
    # Convert each row (tc) to a dictionary
    data = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        row_data = {}
        for i, value in enumerate(row):
            if worksheet.cell(row=1, column=i+1).data_type == 's':
                row_data[headers[i]] = str(value)
            elif worksheet.cell(row=1, column=i+1).data_type == 'n':
                row_data[headers[i]] = float(value)
            elif worksheet.cell(row=1, column=i+1).data_type == 'd':
                row_data[headers[i]] = value.date()
            elif worksheet.cell(row=1, column=i+1).data_type == 'b':
                row_data[headers[i]] = bool(value)
            elif worksheet.cell(row=1, column=i+1).data_type == 'f':
                row_data[headers[i]] = worksheet.cell(row=row[0].row, column=row[0].column)._value
            else:
                row_data[headers[i]] = None
        data.append(row_data)

    process_dict = {process.name: process for process in model.processes.values()}

    for flow in data:
        new_flow = Flow(flow['process_from'], flow['process_to'], flow['composition'])
        model.add_flow(new_flow)
        new_flow.add_to_model(model)
        
        if flow['process_from'] in process_dict:
            process = process_dict[flow['process_from']]
            process.add_output(new_flow)
            
        else:
            print("Process not found, adding: " + flow['process_from'])
            process = Process(flow['process_from'])
            process.add_output(new_flow)
            model.add_process(process)
            # process.add_to_model(model)

        if flow['process_to'] in process_dict:
            process = process_dict[flow['process_to']]
            process.add_input(new_flow)
        else:
            print("Process not found, adding: " + flow['process_to'])
            process = Process(flow['process_from'])
            process.add_input(new_flow)
            model.add_process(process)
            process.add_to_model(model)

    return data


